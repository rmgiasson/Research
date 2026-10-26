#!/usr/bin/env python3
"""
Basal ganglia pathway experiment on the SynthSeg MRI features (exploratory).

STAGE 1 (default) -- feature mapping only (what the most recent email asks for):
    Maps Josh's direct / indirect pathway structures onto the columns of the data file, checks that
    every left/right column exists, and writes basal_ganglia_feature_mapping.xlsx.
    Left and right hemispheres are analysed independently. No model is run.

STAGE 2 (--analysis) -- the exploratory run described in the earlier email, to use after the mapping is reviewed:
    ICV-normalise + standardise the pathway features of ONE cohort, fit DirectLiNGAM separately for
    each hemisphere, draw each estimated
    graph with feature names and edge weights, and compare it with the pathways at the level of
    ANATOMICAL ADJACENCIES first (direction second): which expected connections are recovered, which
    are missing, and which additional connections appear.

Usage:
    python3 basal_ganglia_pathways.py synthseg_HCP.csv                 # stage 1 only
    python3 basal_ganglia_pathways.py synthseg_HCP.csv --analysis      # stages 1 + 2
    python3 basal_ganglia_pathways.py synthseg_HCP.csv --analysis --features direct   # 5 features per hemisphere
    python3 basal_ganglia_pathways.py synthseg_HCP.csv --analysis --bootstrap 100     # optional stability run
The data file may be .csv/.tsv (delimiter detected automatically) or .xlsx.
"""
import argparse
import os
import warnings

import numpy as np
import pandas as pd
import networkx as nx
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

warnings.filterwarnings("ignore")

HEMIS = ("left", "right")
ICV_COL = "total intracranial"
ID_COLS = ["subject"]

# (structure, in direct pathway, in indirect pathway, note)
STRUCTURES = [
    ("cerebral cortex", True, True,
     "Whole-hemisphere cortex volume (no parcellation); separate from the cerebral white matter features."),
    ("caudate", True, True, ""),
    ("putamen", True, True, ""),
    ("pallidum", True, True,
     "One pallidum volume per side, so the two pallidal structures in the indirect pathway cannot be separated."),
    ("thalamus", True, True, ""),
    ("ventral DC", False, True,
     "Composite region; from memory it includes the subthalamic nucleus and substantia nigra among others "
     "(to be confirmed with Josh)."),
]

# Anatomical adjacencies implied by the pathways, within one hemisphere:
# (structure A, structure B, pathway direction A->B or None for a loop, pathway)
EXPECTED = [
    ("cerebral cortex", "caudate", True, "direct + indirect"),
    ("cerebral cortex", "putamen", True, "direct + indirect"),
    ("caudate", "pallidum", True, "direct + indirect"),
    ("putamen", "pallidum", True, "direct + indirect"),
    ("pallidum", "thalamus", True, "direct + indirect"),
    ("thalamus", "cerebral cortex", True, "direct + indirect"),
    ("pallidum", "ventral DC", False, "indirect only (pallidum -> ventral DC -> pallidum loop)"),
]

REASON_OVERRIDES = {
    ICV_COL: "Used only to normalise volumes by head size; not a feature.",
    "subject": "Subject identifier.",
    "csf": "Extra unsided column in this file (not one of the 31 features); not part of either pathway.",
}


def side_of(col):
    for h in HEMIS:
        if col.startswith(h + " "):
            return h
    return "unsided"


def feature(h, s):
    return f"{h} {s}"


def feature_set(kind):
    structs = [s for s, d, i, _ in STRUCTURES if (d if kind == "direct" else i)]
    return [feature(h, s) for h in HEMIS for s in structs]


def load_table(path):
    if os.path.splitext(path)[1].lower() in (".xlsx", ".xls"):
        df = pd.read_excel(path)
    else:
        df = pd.read_csv(path, sep=None, engine="python")   # detects comma or tab
    df.columns = [str(c).strip() for c in df.columns]
    return df


# ======================================================================
# Stage 1: feature mapping workbook
# ======================================================================
def build_mapping(cols):
    rows = []
    for s, direct, indirect, note in STRUCTURES:
        l, r = feature("left", s), feature("right", s)
        rows.append({
            "Pathway structure": s,
            "Left feature": l, "Right feature": r,
            "Left column found in file": "Yes" if l in cols else "NO",
            "Right column found in file": "Yes" if r in cols else "NO",
            "Measured separately left/right": "Yes",
            "In direct pathway": "Yes" if direct else "No",
            "In indirect pathway": "Yes" if indirect else "No",
            "Notes": note,
        })
    return pd.DataFrame(rows)


def build_excluded(cols):
    mapped = {feature(h, s) for h in HEMIS for s, *_ in STRUCTURES}
    rows = []
    for c in cols:
        if c in mapped:
            continue
        if c in REASON_OVERRIDES:
            reason = REASON_OVERRIDES[c]
        elif c.endswith("accumbens area"):
            reason = ("Ventral striatum; could arguably be grouped with caudate/putamen, but Josh's pathways "
                      "do not name it.")
        else:
            reason = "Not part of either pathway."
        rows.append({"Column in file": c, "Side": side_of(c), "Reason not used": reason})
    return pd.DataFrame(rows)


def build_expected():
    rows = []
    for h in HEMIS:
        for a, b, directed, pathway in EXPECTED:
            rows.append({"Hemisphere": h, "Feature A": feature(h, a), "Feature B": feature(h, b),
                         "Pathway direction": f"{feature(h, a)} -> {feature(h, b)}" if directed
                         else "loop (either direction)", "Pathway": pathway})
    return pd.DataFrame(rows)


def write_workbook(path, mapping, excluded, expected, data_file, n_rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        mapping.to_excel(xw, sheet_name="Pathway features", index=False)
        excluded.to_excel(xw, sheet_name="Not in pathways", index=False)
        expected.to_excel(xw, sheet_name="Expected adjacencies", index=False)
        pd.DataFrame({"Item": ["Source file", "Rows in file", "Left/right handling", "Assumptions"],
                      "Detail": [os.path.basename(data_file), n_rows if n_rows > 0 else "0 (header row only; column check)",
                                 "Separate models for left and right hemispheres; no cross-hemisphere edges.",
                                 "Expected adjacencies are ipsilateral only. Pallidum is one feature per side. "
                                 "Ventral DC composition to be confirmed."]}
                     ).to_excel(xw, sheet_name="Notes", index=False)
        for ws in xw.book.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.row == 1:
                        cell.fill = PatternFill("solid", start_color="D9E8FB")
            for k, col in enumerate(ws.columns, start=1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(k)].width = min(max(14, width + 2), 60)
            ws.freeze_panes = "A2"


# ======================================================================
# Stage 2: DirectLiNGAM on the pathway features
# ======================================================================
def normalise_and_standardise(df, features, name):
    from sklearn.preprocessing import StandardScaler
    need = [ICV_COL] + features
    missing = [c for c in need if c not in df.columns]
    if missing:
        raise ValueError(f"{name} is missing columns: {missing}")
    d = df[need].apply(pd.to_numeric, errors="coerce")
    before = len(d)
    d = d.dropna()
    d = d[d[ICV_COL] > 0]
    print(f"  subjects: {before} -> {len(d)} after removing NaN / invalid ICV")
    norm = d[features].div(d[ICV_COL], axis=0)
    if not np.isfinite(norm.values).all():
        raise ValueError("non-finite values after ICV normalisation")
    return StandardScaler().fit_transform(norm.values).astype(np.float64)


def fit_directlingam(X, seed=42):
    import lingam
    m = lingam.DirectLiNGAM(random_state=seed)
    m.fit(X)
    W = m.adjacency_matrix_.T.copy()       # lingam stores [target, source]; we use [source, target]
    np.fill_diagonal(W, 0.0)
    return W, list(m.causal_order_)


def bootstrap_directlingam(X, n_boot, threshold=0.5, seed=42):
    import lingam
    rng = np.random.RandomState(seed)
    n, d = X.shape
    wsum, cnt, ok = np.zeros((d, d)), np.zeros((d, d)), 0
    for b in range(n_boot):
        idx = rng.choice(n, n, replace=True)
        try:
            m = lingam.DirectLiNGAM(random_state=seed + b)
            m.fit(X[idx])
            a = m.adjacency_matrix_.T.copy()
            np.fill_diagonal(a, 0.0)
            wsum += a
            cnt += (np.abs(a) > 0)
            ok += 1
        except Exception as e:
            print(f"  bootstrap {b + 1} failed: {e}")
    probs = cnt / max(ok, 1)
    mean_w = np.divide(wsum, cnt, out=np.zeros_like(wsum), where=cnt > 0)
    return np.where(probs >= threshold, mean_w, 0.0), probs


def compare(W, features):
    """Adjacency-first comparison. Returns (expected_table, additional_table, summary dict)."""
    est = {(s, t): W[i, j] for i, s in enumerate(features) for j, t in enumerate(features)
           if i != j and abs(W[i, j]) > 1e-8}
    exp_rows, exp_pairs = [], set()
    for h in HEMIS:
        for a, b, directed, pathway in EXPECTED:
            fa, fb = feature(h, a), feature(h, b)
            if fa not in features or fb not in features:
                continue
            exp_pairs.add(frozenset((fa, fb)))
            ab, ba = est.get((fa, fb)), est.get((fb, fa))
            if ab is None and ba is None:
                status, direction, w, match = "MISSING", "", np.nan, ""
            else:
                status = "RECOVERED"
                src, tgt = (fa, fb) if ab is not None else (fb, fa)
                direction, w = f"{src} -> {tgt}", est[(src, tgt)]
                match = "Yes" if (not directed or (src, tgt) == (fa, fb)) else "No (reversed)"
            exp_rows.append({"hemisphere": h, "expected adjacency": f"{fa} -- {fb}", "pathway": pathway,
                             "pathway direction": f"{fa} -> {fb}" if directed else "loop (either)",
                             "status": status, "estimated edge": direction, "weight": w,
                             "direction matches pathway": match})
    add_rows = []
    for (s, t), w in est.items():
        if frozenset((s, t)) in exp_pairs:
            continue
        hs, ht = side_of(s), side_of(t)
        if hs != ht:
            kind = ("cross-hemisphere, homologous structure" if s.split(" ", 1)[1] == t.split(" ", 1)[1]
                    else "cross-hemisphere, different structures")
        else:
            kind = "same hemisphere, not in pathway"
        add_rows.append({"estimated edge": f"{s} -> {t}", "weight": w, "type": kind})
    exp_df = pd.DataFrame(exp_rows)
    add_df = pd.DataFrame(add_rows, columns=["estimated edge", "weight", "type"])
    rec = exp_df[exp_df.status == "RECOVERED"]
    summary = {"expected adjacencies": len(exp_df), "recovered": len(rec),
               "missing": int((exp_df.status == "MISSING").sum()),
               "recovered with pathway direction": int((rec["direction matches pathway"] == "Yes").sum()),
               "recovered but reversed": int((rec["direction matches pathway"] == "No (reversed)").sum()),
               "additional edges": len(add_df), "total estimated edges": len(est)}
    return exp_df, add_df, summary


SHORT = {"cerebral cortex": "Cortex", "caudate": "Caudate", "putamen": "Putamen",
         "pallidum": "Pallidum", "thalamus": "Thalamus", "ventral DC": "Ventral DC"}
LOCAL = {"cerebral cortex": (1.0, 4.0), "caudate": (0.0, 3.0), "putamen": (2.0, 3.0),
         "pallidum": (1.0, 1.8), "thalamus": (3.0, 2.6), "ventral DC": (2.3, 0.5)}
OFFSET = {"left": 0.0, "right": 5.5}


def layout(features):
    pos, lab = {}, {}
    for f in features:
        h, s = f.split(" ", 1)
        pos[f] = (LOCAL[s][0] + OFFSET[h], LOCAL[s][1])
        lab[f] = f"{SHORT[s]}\n({h[0].upper()})"
    return pos, lab


def draw_estimate(W, features, exp_df, title, path, probs=None):
    pos, lab = layout(features)
    idx = {f: i for i, f in enumerate(features)}
    G = nx.DiGraph(); G.add_nodes_from(features)
    expected_dir = {}
    for _, r in exp_df.iterrows():
        a, b = r["expected adjacency"].split(" -- ")
        expected_dir[frozenset((a, b))] = (a, b, r["pathway direction"] != "loop (either)")
    groups = {"ok": [], "rev": [], "extra": []}
    labels = {}
    for s in features:
        for t in features:
            if s == t or abs(W[idx[s], idx[t]]) < 1e-8:
                continue
            key = frozenset((s, t))
            if key in expected_dir:
                a, b, directed = expected_dir[key]
                groups["ok" if (not directed or (s, t) == (a, b)) else "rev"].append((s, t))
            else:
                groups["extra"].append((s, t))
            labels[(s, t)] = f"{W[idx[s], idx[t]]:.2f}" + (f"\np={probs[idx[s], idx[t]]:.2f}" if probs is not None else "")
    missing = [tuple(k) for k, v in expected_dir.items()
               if not (abs(W[idx[v[0]], idx[v[1]]]) > 1e-8 or abs(W[idx[v[1]], idx[v[0]]]) > 1e-8)]
    fig, ax = plt.subplots(figsize=(14, 7))
    nx.draw_networkx_nodes(G, pos, node_color="#d9e8fb", edgecolors="black", node_size=2600, ax=ax)
    nx.draw_networkx_labels(G, pos, labels=lab, font_size=8, ax=ax)
    nx.draw_networkx_edges(nx.Graph(missing), pos, edge_color="lightgray", style="dotted", width=2, ax=ax)
    for key, color in (("extra", "red"), ("rev", "orange"), ("ok", "green")):
        if groups[key]:
            nx.draw_networkx_edges(G, pos, edgelist=groups[key], edge_color=color, width=2, arrows=True,
                                   arrowstyle="-|>", arrowsize=16, node_size=2600, ax=ax)
    nx.draw_networkx_edge_labels(G, pos, edge_labels=labels, font_size=7, label_pos=0.4,
                                 bbox=dict(boxstyle="round,pad=0.1", fc="white", ec="none", alpha=0.8), ax=ax)
    ax.legend(handles=[Line2D([0], [0], color="green", lw=2, label="Expected adjacency, pathway direction"),
                       Line2D([0], [0], color="orange", lw=2, label="Expected adjacency, reversed"),
                       Line2D([0], [0], color="red", lw=2, label="Additional (not in pathways)"),
                       Line2D([0], [0], color="lightgray", lw=2, ls=":", label="Expected adjacency missing")],
              loc="upper center", bbox_to_anchor=(0.5, -0.01), ncol=4, frameon=False, fontsize=9)
    ax.set_title(title); ax.axis("off")
    fig.tight_layout(); fig.savefig(path, dpi=200, bbox_inches="tight"); plt.close(fig)


def run_hemisphere_analysis(df, args, out, hemisphere):
    features = [f for f in feature_set(args.features) if side_of(f) == hemisphere]
    print(f"\n=== DirectLiNGAM: {hemisphere}, {len(features)} '{args.features}' features ({os.path.basename(args.data)}) ===")
    X = normalise_and_standardise(df, features, f"{hemisphere} cohort")
    print(f"  input: {X.shape[0]} samples x {X.shape[1]} variables")
    runs = {}
    W, order = fit_directlingam(X)
    runs["single_fit"] = (W, None)
    pd.DataFrame({"causal_order": [features[i] for i in order]}).to_csv(os.path.join(out, f"causal_order_{args.features}_{hemisphere}.csv"), index=False)
    if args.bootstrap > 0:
        print(f"  bootstrap: {args.bootstrap} resamples")
        runs["bootstrap"] = bootstrap_directlingam(X, args.bootstrap)
    for mode, (Wm, probs) in runs.items():
        tag = f"{args.features}_{hemisphere}_{mode}"
        pd.DataFrame(Wm, index=features, columns=features).to_csv(os.path.join(out, f"adjacency_{tag}.csv"))
        if probs is not None:
            pd.DataFrame(probs, index=features, columns=features).to_csv(os.path.join(out, f"edge_probabilities_{tag}.csv"))
        exp_df, add_df, summary = compare(Wm, features)
        exp_df.to_csv(os.path.join(out, f"expected_connections_{tag}.csv"), index=False)
        add_df.to_csv(os.path.join(out, f"additional_connections_{tag}.csv"), index=False)
        pd.Series(summary).to_csv(os.path.join(out, f"comparison_summary_{tag}.csv"), header=["value"])
        draw_estimate(Wm, features, exp_df, f"DirectLiNGAM ({mode}), {hemisphere} hemisphere, {args.features} features, "
                      f"n={X.shape[0]}  [source -> target, edge labels = weights]",
                      os.path.join(out, f"graph_{tag}.png"), probs)
        print(f"\n--- {mode} ---")
        print("Summary:", summary)
        print("\nExpected connections:")
        print(exp_df[["expected adjacency", "status", "estimated edge", "weight", "direction matches pathway"]]
              .to_string(index=False))
        print("\nAdditional connections:")
        print(add_df.to_string(index=False) if len(add_df) else "  none")


def run_analysis(df, args, out):
    for hemisphere in HEMIS:
        run_hemisphere_analysis(df, args, out, hemisphere)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("data", help="HCP (or other cohort) file: .csv / .tsv / .xlsx")
    ap.add_argument("--analysis", action="store_true", help="also run the DirectLiNGAM stage")
    ap.add_argument("--features", choices=["direct", "indirect"], default="indirect",
                    help="per hemisphere: 'indirect' = 6 features incl. ventral DC (default); 'direct' = 5 features")
    ap.add_argument("--bootstrap", type=int, default=0, help="optional number of bootstrap resamples")
    ap.add_argument("--outdir", default="basal_ganglia_outputs")
    args = ap.parse_args()

    os.makedirs(args.outdir, exist_ok=True)
    df = load_table(args.data)
    cols = list(df.columns)
    print(f"Loaded {os.path.basename(args.data)}: {len(df)} rows, {len(cols)} columns")

    mapping, excluded, expected = build_mapping(cols), build_excluded(cols), build_expected()
    path = os.path.join(args.outdir, "basal_ganglia_feature_mapping.xlsx")
    write_workbook(path, mapping, excluded, expected, args.data, len(df))
    missing = [c for c in ("Left column found in file", "Right column found in file") for v in mapping[c] if v != "Yes"]
    print(f"Saved {path}")
    print(mapping[["Pathway structure", "Left feature", "Right feature", "Left column found in file",
                   "Right column found in file"]].to_string(index=False))
    print(f"ICV column '{ICV_COL}' present: {ICV_COL in cols}")
    if missing:
        print("WARNING: some pathway columns are missing from the file; stage 2 cannot run.")
    if args.analysis:
        if missing or ICV_COL not in cols:
            raise SystemExit("Cannot run analysis: required columns missing.")
        run_analysis(df, args, args.outdir)
    else:
        print("\n(Stage 1 only. Re-run with --analysis to run DirectLiNGAM after reviewing the mapping.)")


if __name__ == "__main__":
    main()
